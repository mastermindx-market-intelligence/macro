"""Official country inputs for the existing international-macro collector.

BIS observations and attributed GDP forecasts remain separate. Parsers never
replace a missing observation with zero or treat a forecast as realized growth.
Network acquisition is injected; the render path performs no network requests.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from hashlib import sha256
import io
import math
import re
from copy import deepcopy
import json
from pathlib import Path
from datetime import timezone
from lib.intl_forecast_vintages import merge_publications, select_publications
from lib.intl_policy_evidence import merge_policy_decisions, merge_policy_reviews
from collectors.intl_policy_sources import collect_registered_policy
from collectors.intl_weo import collect_registered_weo
from zipfile import ZipFile, BadZipFile

SCHEMA = "intl_country_inputs.v1"
BIS_CPI = "https://data.bis.org/static/bulk/WS_LONG_CPI_csv_flat.zip"
BIS_POLICY = "https://data.bis.org/static/bulk/WS_CBPOL_csv_flat.zip"
IMF_XLSX = "https://www.imf.org/-/media/files/publications/weo/2026/update/july/english/data/weojul2026.xlsx"
IMF_PAGE = "https://www.imf.org/en/publications/weo/issues/2026/07/08/world-economic-outlook-update-july-2026"
IMF_PUBLISHED = "2026-07-08"
TW_PAGE = "https://eng.dgbas.gov.tw/News_Content.aspx?n=4438&s=236587"
TW_PUBLISHED = "2026-08-14"
AREAS = {"JP": "JP", "KR": "KR", "TW": "TW", "IN": "IN", "AU": "AU", "GB": "GB", "XM": "EZ"}
COUNTRY_LABELS = {"Japan": "JP", "Korea": "KR", "India": "IN", "Australia": "AU",
                  "United Kingdom": "GB", "Euro Area": "EZ"}


def _code(value):
    return str(value or "").split(":", 1)[0].strip()


def _finite(value):
    if type(value) is bool:
        return None
    try:
        number = float(value)
    except (ValueError, TypeError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def _month(value):
    if type(value) is not str or not re.fullmatch(r"\d{4}-\d{2}", value):
        raise ValueError("invalid_observation_period")
    date.fromisoformat(value + "-01")
    return value


def parse_bis(blob, kind, *, retrieved_at):
    """Read public, monthly, source-defined observations; bound ZIP expansion."""
    if kind not in ("inflation", "policy") or type(blob) is not bytes or len(blob) > 16_000_000:
        raise ValueError("invalid_bis_input")
    dataset = "WS_LONG_CPI" if kind == "inflation" else "WS_CBPOL"
    expected_unit = "771" if kind == "inflation" else "368"
    cutoff = date.fromisoformat(retrieved_at[:10])
    series = {}
    with ZipFile(io.BytesIO(blob)) as archive:
        entries = [x for x in archive.infolist() if x.filename.endswith(".csv")]
        # The actual policy archive is475.9MB because it repeats series metadata.
        # Stream it; never extract files or allocate its entire expanded contents.
        if len(entries) != 1 or entries[0].filename != dataset + "_csv_flat.csv" or entries[0].file_size > 600_000_000:
            raise ValueError("unexpected_bis_archive")
        with archive.open(entries[0]) as handle:
            rows = csv.DictReader(io.TextIOWrapper(handle, encoding="utf-8-sig"))
            required = {"STRUCTURE_ID", "FREQ:Frequency", "REF_AREA:Reference area",
                        "TIME_PERIOD:Time period or range", "OBS_VALUE:Observation Value",
                        "UNIT_MEASURE:Unit of measure", "UNIT_MULT:Unit Multiplier",
                        "OBS_CONF:Observation confidentiality", "OBS_STATUS:Observation Status"}
            if not required <= set(rows.fieldnames or []):
                raise ValueError("unexpected_bis_columns")
            for row in rows:
                area = _code(row["REF_AREA:Reference area"])
                if area not in AREAS or _code(row["FREQ:Frequency"]) != "M":
                    continue
                if not row["STRUCTURE_ID"].startswith("BIS:" + dataset + "(1.0):"):
                    raise ValueError("unexpected_bis_dataset")
                if _code(row["UNIT_MEASURE:Unit of measure"]) != expected_unit:
                    continue
                if _code(row["UNIT_MULT:Unit Multiplier"]) != "0":
                    raise ValueError("unexpected_bis_multiplier")
                if _code(row["OBS_CONF:Observation confidentiality"]) != "F":
                    continue
                status = _code(row["OBS_STATUS:Observation Status"])
                if status not in ("A", "P"):
                    continue
                period = _month(row["TIME_PERIOD:Time period or range"])
                # Month-end data for an incomplete future month is not an observation.
                if period >= cutoff.strftime("%Y-%m"):
                    continue
                value = _finite(row["OBS_VALUE:Observation Value"])
                if value is None:
                    continue
                cc = AREAS[area]
                points = series.setdefault(cc, {})
                if period in points and points[period]["value"] != value:
                    raise ValueError("conflicting_bis_observation")
                points[period] = {
                    "value": value, "status": "provisional" if status == "P" else "observed",
                    "definition": row.get("COMPILATION:Compilation") or row.get("COVERAGE:Coverage") or "",
                }
    observations = {}
    url = BIS_CPI if kind == "inflation" else BIS_POLICY
    for cc, points in series.items():
        last = max(points)
        item = points[last]
        year, month = map(int, last.split("-"))
        back = year * 12 + month - 1 - 3
        prior = f"{back // 12:04d}-{back % 12 + 1:02d}"
        observations[cc] = {
            "metric": kind, "value": item["value"], "period": last,
            "unit": "% y/y" if kind == "inflation" else "%",
            "provider": "Bank for International Settlements",
            "source_url": "https://data.bis.org/topics/CPI" if kind == "inflation" else "https://data.bis.org/topics/CBPOL",
            "dataset_url": url, "series_id": f"M.{next(k for k,v in AREAS.items() if v == cc)}" + (".771" if kind == "inflation" else ""),
            "status": item["status"], "frequency": "monthly",
            "definition": item["definition"], "retrieved_at": retrieved_at,
            "change_3m_pp": round(item["value"] - points[prior]["value"], 6) if prior in points else None,
            "source_sha256": sha256(blob).hexdigest(),
        }
    return {"observations": observations, "series": series}


def parse_imf_workbook(blob, *, retrieved_at):
    """Validate the observed July edition; preserve the fiscal-year exception."""
    from openpyxl import load_workbook
    if type(blob) is not bytes or len(blob) > 4_000_000:
        raise ValueError("invalid_imf_input")
    workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        cover = workbook["WEO Update July 2026"]
        published = cover.cell(3, 9).value
        if not isinstance(published, datetime) or published.date().isoformat() != IMF_PUBLISHED:
            raise ValueError("unexpected_imf_vintage")
        output = {}
        for sheet_name in ("Table 1.", "Annex Table 1."):
            rows = list(workbook[sheet_name].values)
            header = next((r for r in rows[:12] if tuple(r[1:5]) == (2024, 2025, 2026, 2027)), None)
            if header is None or not any("Projections" in r for r in rows[:12]):
                raise ValueError("unexpected_imf_periods")
            for row in rows:
                label = re.sub(r"\s+\d+/", "", str(row[0] or "")).strip()
                cc = COUNTRY_LABELS.get(label)
                if cc is None:
                    continue
                values = {str(year): _finite(row[index]) for index, year in enumerate((2024, 2025, 2026, 2027), 1)}
                if values["2026"] is None or values["2027"] is None:
                    raise ValueError("missing_imf_projection")
                result = {
                    "cc": cc, "kind": "sourced_forecast", "provider": "IMF",
                    "source_url": IMF_PAGE, "dataset_url": IMF_XLSX,
                    "published": IMF_PUBLISHED, "retrieved_at": retrieved_at,
                    "unit": "percent", "period_basis": "fiscal_year" if cc == "IN" else "calendar_year",
                    "values": values, "revision_current_pp": _finite(row[6]), "revision_from": "2026-04",
                    "source_sha256": sha256(blob).hexdigest(),
                }
                if cc in output and output[cc]["values"] != values:
                    raise ValueError("conflicting_imf_projection")
                output[cc] = result
        if set(output) != set(COUNTRY_LABELS.values()):
            raise ValueError("incomplete_imf_country_table")
        return output
    finally:        workbook.close()


def parse_taiwan_release(html, *, retrieved_at):
    """Use one explicit official release, never a search snippet or a later date."""
    from html.parser import HTMLParser
    class Text(HTMLParser):
        def __init__(self):
            super().__init__()
            self.parts = []
        def handle_data(self, data):
            self.parts.append(data)
    if type(html) is not str or len(html) > 2_000_000:
        raise ValueError("invalid_taiwan_release")
    parser = Text()
    parser.feed(html)
    text = " ".join(" ".join(parser.parts).split())
    if "August 14, 2026" not in text:
        raise ValueError("unexpected_taiwan_vintage")
    pattern = r"real GDP is predicted to grow\s+([0-9.]+)%\s+and\s+([0-9.]+)%\s+in\s+2026\s+and\s+2027"
    matches = re.findall(pattern, text, flags=re.I)
    if not matches or len(set(matches)) != 1:
        raise ValueError("unresolved_taiwan_projection")
    first, second = map(float, matches[0])
    if not all(math.isfinite(v) and 0 <= v < 50 for v in (first, second)):
        raise ValueError("invalid_taiwan_projection")
    return {"TW": {
        "cc": "TW", "kind": "sourced_forecast", "provider": "DGBAS",
        "source_url": TW_PAGE, "published": TW_PUBLISHED, "retrieved_at": retrieved_at,
        "unit": "percent", "period_basis": "calendar_year",
        "values": {"2026": first, "2027": second},
        "revision_current_pp": None,
        "source_sha256": sha256(html.encode("utf-8")).hexdigest(),
    }}


def collect_country_inputs(fetch_bytes, *, retrieved_at, releases=None):
    """Partial independent sources, with explicit errors and no manufactured data."""
    official, forecasts, series, errors, sources = {}, {}, {}, [], []
    for kind, url in (("inflation", BIS_CPI), ("policy", BIS_POLICY)):
        try:
            blob = fetch_bytes(url, max_bytes=16_000_000)
            parsed = parse_bis(blob, kind, retrieved_at=retrieved_at)
            for cc, item in parsed["observations"].items():
                official.setdefault(cc, {})[kind] = item
            if kind == "inflation":
                series = parsed["series"]
            sources.append({"url": url, "sha256": sha256(blob).hexdigest(), "status": "read", "retrieved_at": retrieved_at})
        except (ValueError, OSError, KeyError, csv.Error, BadZipFile) as exc:
            errors.append({"source": kind, "error": type(exc).__name__})
    for key, url, parser in (("imf", IMF_XLSX, parse_imf_workbook), ("taiwan", TW_PAGE, parse_taiwan_release)):
        try:
            blob = fetch_bytes(url, max_bytes=4_000_000)
            argument = blob.decode("utf-8") if key == "taiwan" else blob
            forecasts.update(parser(argument, retrieved_at=retrieved_at))
            sources.append({"url": url, "sha256": sha256(blob).hexdigest(), "status": "read", "retrieved_at": retrieved_at})
        except (ValueError, OSError, KeyError, UnicodeError, BadZipFile) as exc:
            errors.append({"source": key, "error": type(exc).__name__})
    if releases is None:
        try:
            catalogue = json.loads((Path(__file__).resolve().parents[1] / "config/intl_country_releases.json").read_text())
            if catalogue.get("schema") != "intl_country_release_catalogue.v1":
                raise ValueError("invalid_release_catalogue_schema")
            releases = catalogue["releases"]
        except (OSError, ValueError, KeyError) as exc:
            releases = []
            errors.append({"source": "WEO release catalogue", "error": type(exc).__name__})
    current = collect_registered_weo(fetch_bytes, releases=releases, retrieved_at=retrieved_at)
    policy = collect_registered_policy(fetch_bytes, retrieved_at=retrieved_at)
    return {"schema": SCHEMA, "retrieved_at": retrieved_at,
            "official": official, "forecasts": forecasts,
            "forecast_archive": current["forecast_archive"] + policy["forecast_archive"], "cpi_series": series,
            "policy_decisions": policy["policy_decisions"], "policy_reviews": policy["policy_reviews"],
            "sources": sources + current["sources"] + policy["sources"], "errors": errors + current["errors"] + policy["errors"]}

def _mapping(value):
    return value if type(value) is dict else {}


def _source_clock(value, *, monthly=False):
    if type(value) is not str:
        return None
    try:
        return _month(value) if monthly else date.fromisoformat(value).isoformat()
    except ValueError:
        return None


def _valid_incoming(value, *, cc, metric=None):
    """Reject malformed refresh entries before they replace a last good value."""
    if type(value) is not dict or not str(value.get("source_url", "")).startswith("https://"):
        return False
    if not isinstance(value.get("provider"), str) or not value["provider"].strip():
        return False
    if metric is not None:
        unit = "% y/y" if metric == "inflation" else "%"
        return (value.get("metric") == metric and value.get("unit") == unit
                and value.get("status") in ("observed", "provisional")
                and type(value.get("value")) in (int, float)
                and _finite(value.get("value")) is not None
                and _source_clock(value.get("period"), monthly=True) is not None)
    values = _mapping(value.get("values"))
    return (value.get("cc") == cc and value.get("kind") == "sourced_forecast"
            and value.get("unit") == "percent"
            and value.get("period_basis") in ("calendar_year", "fiscal_year")
            and _source_clock(value.get("published")) is not None
            and len(values) >= 2
            and all(type(v) in (int, float) and _finite(v) is not None and -100 < v < 100
                    for k, v in values.items() if str(k).isdigit())
            and sum(str(k).isdigit() for k in values) >= 2)


def merge_country_inputs(previous, incoming, *, today=None, preserve_custody=False):
    """Retain publications, then select each country; never replace by arrival."""
    today = today or date.today()
    result = deepcopy(previous) if type(previous) is dict and previous.get("schema") == SCHEMA else {}
    incoming = incoming if type(incoming) is dict and incoming.get("schema") == SCHEMA else {}
    result["schema"] = SCHEMA
    result["official"] = {cc: deepcopy(value) for cc, value in _mapping(result.get("official")).items()
                          if cc in AREAS.values() and type(value) is dict}
    for cc, values in _mapping(incoming.get("official")).items():
        if cc not in AREAS.values() or type(values) is not dict:
            continue
        existing = result["official"].setdefault(cc, {})
        for key, observation in values.items():
            if key not in ("inflation", "policy") or not _valid_incoming(observation, cc=cc, metric=key):
                continue
            current = _mapping(existing.get(key))
            current_clock = _source_clock(current.get("period"), monthly=True) or ""
            if observation["period"] >= current_clock:
                existing[key] = deepcopy(observation)
    def publications(envelope):
        records = deepcopy(envelope.get("forecast_archive")) if type(envelope.get("forecast_archive")) is list else []
        records.extend(value for cc, value in _mapping(envelope.get("forecasts")).items()
                       if cc in AREAS.values() and type(value) is dict)
        return records
    observed_at = datetime.now(timezone.utc).isoformat()
    archive, diagnostics = merge_publications(publications(result), publications(incoming), observed_at=observed_at)
    result["forecast_archive"] = archive
    result["policy_decisions"] = merge_policy_decisions(result.get("policy_decisions"), incoming.get("policy_decisions"),
                                                       observed_at=observed_at, preserve_custody=preserve_custody)
    result["policy_reviews"] = merge_policy_reviews(result.get("policy_reviews"), incoming.get("policy_reviews"),
                                                   observed_at=observed_at, preserve_custody=preserve_custody)
    result["forecast_diagnostics"] = diagnostics
    result["forecasts"] = {}
    for cc in set(AREAS.values()):
        eligible = select_publications(archive, cc, today=today)
        if eligible:
            result["forecasts"][cc] = deepcopy(eligible[0])
    result["cpi_series"] = deepcopy(_mapping(result.get("cpi_series")))
    for cc, series in _mapping(incoming.get("cpi_series")).items():
        if cc in AREAS.values() and type(series) is dict and series:
            result["cpi_series"][cc] = deepcopy(series)
    # Each receipt belongs to one artifact. Replaying a partial cache must not
    # erase a newer source, nor let a page build refresh another source's clock.
    receipts = {}
    for envelope in (result, incoming):
        for item in envelope.get("sources", []) if type(envelope.get("sources")) is list else []:
            if type(item) is not dict:
                continue
            url = item.get("dataset_url") or item.get("url") or item.get("source_url")
            digest = item.get("sha256")
            if not isinstance(url, str) or not url.startswith("https://") or not isinstance(digest, str):
                continue
            identity = (url, digest)
            previous_receipt = receipts.get(identity, {})
            combined = {**deepcopy(previous_receipt), **deepcopy(item)}
            clocks = [x.get("retrieved_at") for x in (previous_receipt, item)
                      if isinstance(x.get("retrieved_at"), str)]
            if clocks:
                combined["retrieved_at"] = max(clocks)
            receipts[identity] = combined
    if len(receipts) > 100:
        raise ValueError("country_source_receipt_capacity")
    result["sources"] = [receipts[key] for key in sorted(receipts)]
    clocks = [envelope.get("retrieved_at") for envelope in (result, incoming)
              if isinstance(envelope.get("retrieved_at"), str)]
    if clocks:
        result["retrieved_at"] = max(clocks)
    if "errors" in incoming and type(incoming["errors"]) is list:
        result["errors"] = deepcopy(incoming["errors"])
        result["errors_observed_at"] = incoming.get("errors_observed_at", incoming.get("retrieved_at"))
    # Match the reader's serialized input budget; never publish an unreadable
    # archive or silently prune accepted history to make it fit.
    if len(json.dumps(result, ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")) > 1_500_000:
        raise ValueError("country_input_byte_capacity")
    return result
